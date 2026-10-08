from __future__ import annotations

import json
import shutil
import math
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
SOURCE_BOOK = ROOT / "设计表.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "resource_catalog"
OUTPUT_BOOK = OUTPUT_DIR / "现有资源设计表.xlsx"

ELEMENTS = {"metal": "金", "wood": "木", "water": "水", "fire": "火", "earth": "土"}
RARITIES = {"common": "普通", "uncommon": "进阶", "rare": "稀有", "epic": "珍稀"}
UNIT_TYPES = {"beast": "妖兽", "cultivator": "修士", "construct": "傀儡"}
TARGETS = {"self": "自身", "opponent": "敌方", "ally": "友军", "all_opponents": "全体敌方"}


# Excel-only design candidates. These are deliberately kept out of runtime JSON until reviewed.
PROPOSED_SPELLS = [
    ["piercing_rainbow_qi", "贯虹剑气", "金", 1, "普通", "金×2", "造成7点金伤害；目标有护盾时额外造成4点伤害。", "破盾/稳定输出", "普通战、商店"],
    ["hidden_edge_form", "藏锋式", "金", 1, "普通", "金×2", "获得6点护盾与1层锋锐。", "攻防蓄势", "普通战、修炼"],
    ["sword_echo", "剑鸣回响", "金", 2, "进阶", "金×3", "造成9点伤害；本回合每施展过一个金行神通，伤害+2。", "同回合连招", "精英、商店"],
    ["ten_thousand_edges", "万锋归一", "金", 3, "稀有", "金×5", "造成12点伤害，消耗全部锋锐；每层再造成3点伤害。", "锋锐终结", "Boss、稀有机缘"],
    ["root_binding_art", "盘根缚", "木", 1, "普通", "木×2", "造成5点木伤害，并使目标下次攻击降低3点。", "减攻/生存", "普通战、商店"],
    ["green_pulse", "青脉回息", "木", 1, "普通", "木×2", "恢复4点生命并获得1层生机。", "恢复蓄势", "普通战、休息"],
    ["vine_reversal", "回藤返刺", "木", 2, "进阶", "木×3", "获得7点护盾；本回合护盾被击破时，对攻击者造成8点伤害。", "反击窗口", "精英、商店"],
    ["spring_returns", "枯荣一春", "木", 3, "稀有", "木×5", "恢复8点生命并对全体敌人造成8点伤害；消耗全部生机，每层使两项数值+1。", "生机终结", "Boss、稀有机缘"],
    ["still_water_mirror", "止水镜", "水", 1, "普通", "水×2", "获得5点护盾与1层回流。", "防御/回流", "普通战、商店"],
    ["cold_current_bind", "寒流束脉", "水", 1, "普通", "水×2", "造成6点水伤害，使目标下次强化效果降低一半。", "压制强化", "普通战、事件"],
    ["wave_after_wave", "叠浪诀", "水", 2, "进阶", "水×3", "造成7点伤害并返还1点水灵气；若本回合已施展水行神通，再重复一次。", "连招/续航", "精英、商店"],
    ["sea_in_one_breath", "沧海一息", "水", 3, "稀有", "水×5", "对全体造成10点伤害；消耗全部回流，每层返还1点随机灵气。", "群攻终结", "Boss、稀有机缘"],
    ["heart_burning_flame", "焚心火", "火", 1, "普通", "火×2", "对敌人造成10点火伤害，自身受到2点伤害。", "高效风险输出", "普通战、商店"],
    ["furnace_guard", "炉火护脉", "火", 1, "普通", "火×2", "获得5点护盾与1层火种。", "攻防蓄势", "普通战、修炼"],
    ["chasing_flame", "逐焰", "火", 2, "进阶", "火×3", "造成8点伤害；若目标本回合已受过火伤害，再造成6点。", "连续点燃", "精英、商店"],
    ["threefold_true_fire", "三叠真火", "火", 3, "稀有", "火×5", "连续造成3次5点伤害；消耗1层火种后每次伤害+3。", "火种终结", "Boss、稀有机缘"],
    ["bearing_mountain_form", "负岳式", "土", 1, "普通", "土×2", "获得8点护盾与1层地势。", "稳定防御", "普通战、商店"],
    ["earth_pulse", "地脉震", "土", 1, "普通", "土×2", "对全体敌人造成5点土伤害；自身有护盾时伤害+2。", "群攻/盾联动", "普通战、事件"],
    ["stone_skin_reversal", "石衣卸力", "土", 2, "进阶", "土×3", "获得10点护盾；本回合首次受到的生命伤害转为地势。", "承伤蓄势", "精英、商店"],
    ["immovable_land", "不动山河", "土", 3, "稀有", "土×5", "获得16点护盾；消耗全部地势，每层使本回合护盾不衰减并追加2点护盾。", "地势终结", "Boss、稀有机缘"],
]

PROPOSED_RELICS = [
    ["edge_tempering_box", "淬锋匣", "金", "每回合施展第二个金行神通后，下一个金行神通少消耗1点金灵气。", "回合内计数/费用修正", "精英、宝箱"],
    ["sword_shadow_tassel", "剑影穗", "金", "对有护盾的敌人造成伤害后，获得1层锋锐；每回合限一次。", "条件触发/锋锐", "精英、商店"],
    ["broken_bronze_bell", "裂纹铜钟", "金", "自身护盾被完全击破时，对全体敌人造成5点金伤害。", "破盾反击", "Boss、稀有宝箱"],
    ["wither_bloom_branch", "枯荣枝", "木", "生命已满时获得的治疗转化为等量护盾。", "治疗转化", "精英、宝箱"],
    ["spring_thunder_muyu", "春雷木鱼", "木", "每施展三个木行神通，获得1层生机并恢复2点生命。", "跨回合计数", "精英、商店"],
    ["seed_of_return", "归根籽", "木", "生命首次低于一半时，消耗全部生机；每层恢复2点生命。", "濒危自动触发", "Boss、稀有宝箱"],
    ["tide_mark_pearl", "潮痕珠", "水", "回合结束时，至多1点未消耗水灵气不逸散，转化为1层回流。", "资源转状态", "精英、宝箱"],
    ["cold_mirror", "玄冰照影", "水", "每场战斗首次受到伤害时将其减半，并在下一回合少凝聚1点灵气。", "延迟代价", "精英、商店"],
    ["river_chart_fragment", "河图残页", "水", "每回合第一次返还灵气时，额外返还1点同属性灵气。", "返还放大", "Boss、稀有宝箱"],
    ["ash_lamp", "照骨灰灯", "火", "火种被消耗时，对全体敌人造成2点火伤害。", "状态消耗触发", "精英、宝箱"],
    ["fire_washed_silk", "火浣纱", "火", "每回合第一次消耗火灵气时，获得等同当前火种层数的护盾。", "资源消耗触发", "精英、商店"],
    ["three_legged_furnace", "三足火炉", "火", "单回合累计消耗5点火灵气后获得1层火种；每回合限一次。", "回合内计数", "Boss、稀有宝箱"],
    ["mountain_river_inkstone", "山河砚", "土", "护盾被完全击破时获得1层地势。", "破盾蓄势", "精英、宝箱"],
    ["earth_vein_compass", "寻脉盘", "土", "回合结束时若仍有护盾，下回合额外凝聚1点土灵气。", "护盾转资源", "精英、商店"],
    ["burdened_stele", "负岳碑", "土", "每场战斗第一次获得超过10点护盾时，将其中3点转为永久持续至战斗结束的护盾。", "护盾规则改变", "Boss、稀有宝箱"],
]

PROPOSED_ENEMIES = [
    ["mist_reed_crane", "雾苇鹤", "普通", "妖兽", "水", "低", "低", "雾羽啄、敛翼藏雾", "先以薄盾遮身，再寻找空隙啄击。", "教学：防御意图不等于停止进攻。", "mist_reed_crane_rewards"],
    ["stoneback_ape", "负碑石猿", "普通", "妖兽", "土", "中", "中", "碑拳、负石蓄势", "蓄势后重击；受击会提前消耗蓄势换取护盾。", "教学：观察爆发窗口。", "stoneback_ape_rewards"],
    ["golden_thread_spider", "金丝妖蛛", "普通", "妖兽", "金", "低", "中", "金丝针、结茧", "交替攻击与结茧，结茧后下一击更强。", "教学：打断防御循环。", "golden_thread_spider_rewards"],
    ["char_horn_boar", "炭角火豕", "普通", "妖兽", "火", "中", "低", "炭角冲、焦躁刨地", "连续受击会积累焦躁，随后发动冲锋。", "教学：控制输出节奏。", "char_horn_boar_rewards"],
    ["wandering_herbalist", "采药散修", "普通", "修士", "木", "中", "低", "药锄、回春散", "储物袋中携带一次治疗物品，受伤后使用。", "教学：优先击杀与敌人物品。", "wandering_herbalist_rewards"],
    ["grave_moss_wraith", "墓苔阴灵", "普通", "妖兽", "木/阴", "低", "中", "阴苔触、借尸回青", "攻击后缓慢恢复；治疗可被爆发压制。", "教学：持续战与治疗压力。", "grave_moss_wraith_rewards"],
    ["hundred_blade_roamer", "百刃散人", "精英", "修士", "金", "高", "中", "藏剑、百刃落、破甲符", "藏锋数回合后释放剑雨，也会消耗符箓破盾。", "压力：计算防御资源与爆发时机。", "hundred_blade_roamer_rewards"],
    ["marsh_brood_mother", "泽母妖蟾", "精英", "妖兽", "水", "高", "高", "湿毒吐息、鼓腹护卵、唤幼蟾", "在攻击、恢复和召唤之间循环；幼蟾强化本体恢复。", "压力：多目标优先级。", "marsh_brood_mother_rewards"],
    ["furnace_heart_golem", "炉心石傀", "精英", "傀儡", "火/土", "高", "高", "炉甲、泄焰、熔心重锤", "炉温逐轮上升，护盾与伤害交替增强；泄焰后出现输出窗口。", "压力：姿态循环与爆发窗口。", "furnace_heart_golem_rewards"],
    ["five_vein_alchemist", "五脉丹师", "Boss", "修士", "五行", "危险", "危险", "轮转丹火、错脉散、五气炉壁、逆炼归元、炸炉", "按五行轮转改变下一行动；半血后开始逆炼，最终进入炸炉倒计时。", "综合考验：资源规划、状态处理与阶段爆发。", "five_vein_alchemist_rewards"],
]

PROPOSED_ABILITIES = [
    ["mist_peck", "雾羽啄", "造成适中水伤害。", "damage(水,6,opponent)", "mist_reed_crane"],
    ["mist_wing_guard", "敛翼藏雾", "获得护盾；若护盾保留到下回合，下一次攻击伤害提高。", "shield(6,self)+attack_up(3,1回合)", "mist_reed_crane"],
    ["stele_fist", "碑拳", "造成适中土伤害。", "damage(土,7,opponent)", "stoneback_ape"],
    ["bear_stone_charge", "负石蓄势", "获得蓄势；下次碑拳伤害提高。受到较大单次伤害时，改为获得护盾。", "add_status(负石,1,self)", "stoneback_ape"],
    ["golden_silk_needle", "金丝针", "造成两次较弱金伤害。", "damage(金,3)×2", "golden_thread_spider"],
    ["silk_cocoon", "结茧", "获得护盾与1层茧势。", "shield(7,self)+add_status(茧势,1,self)", "golden_thread_spider"],
    ["char_horn_charge", "炭角冲", "消耗全部焦躁，造成火伤害；每层焦躁提高伤害。", "damage(火,6+焦躁×2)", "char_horn_boar"],
    ["restless_paw", "焦躁刨地", "获得强化与1层焦躁。", "attack_up(2,2回合)+add_status(焦躁,1,self)", "char_horn_boar"],
    ["herb_hoe", "药锄", "造成较弱无属性伤害。", "damage(5,opponent)", "wandering_herbalist"],
    ["use_spring_powder", "服回春散", "消耗回春散，恢复生命。", "consume_item(回春散)+heal(10,self)", "wandering_herbalist"],
    ["moss_touch", "阴苔触", "造成木属性伤害并获得1层苔息。", "damage(木,5)+add_status(苔息,1,self)", "grave_moss_wraith"],
    ["borrowed_green", "借尸回青", "消耗苔息恢复生命。", "consume_status(苔息,all)+heal(4+层数×2,self)", "grave_moss_wraith"],
    ["sheathe_hundred_blades", "藏剑", "获得2层藏锋和少量护盾。", "shield(5)+add_status(藏锋,2,self)", "hundred_blade_roamer"],
    ["hundred_blades_fall", "百刃落", "消耗藏锋，对目标连续攻击。", "damage(金,4)×藏锋层数", "hundred_blade_roamer"],
    ["use_armor_break_talisman", "祭破甲符", "消耗符箓，清除目标部分护盾并造成伤害。", "consume_item(破甲符)+damage(金,8)+shield_break(6)", "hundred_blade_roamer"],
    ["damp_poison_breath", "湿毒吐息", "造成水伤害并施加湿毒。", "damage(水,7)+add_status(湿毒,1,opponent)", "marsh_brood_mother"],
    ["brood_guard", "鼓腹护卵", "为自身和幼蟾获得护盾。", "shield(8,self)+shield(5,allies)", "marsh_brood_mother"],
    ["summon_tadpole", "唤幼蟾", "召唤一只泽幼蟾；场上至多两只。", "summon(泽幼蟾)", "marsh_brood_mother"],
    ["furnace_armor", "炉甲", "获得护盾并提高炉温。", "shield(10)+add_status(炉温,1,self)", "furnace_heart_golem"],
    ["vent_flame", "泄焰", "对全体造成火伤害并清空炉温；自身防御短暂下降。", "damage(火,4+炉温×2,all_opponents)+consume_status(炉温,all)+defense_down", "furnace_heart_golem"],
    ["molten_hammer", "熔心重锤", "造成火土伤害；炉温越高伤害越大。", "damage(火/土,8+炉温×3)", "furnace_heart_golem"],
    ["cycle_pill_flame", "轮转丹火", "切换炉中主行，并按当前五行造成不同效果。", "cycle_attribute(金木水火土)", "five_vein_alchemist"],
    ["misaligned_meridian_powder", "错脉散", "使玩家下一次凝聚灵气的属性权重发生偏移。", "add_status(错脉,1,opponent)", "five_vein_alchemist"],
    ["five_qi_furnace_wall", "五气炉壁", "按已轮转过的属性数获得护盾。", "shield(5+已轮转属性×3,self)", "five_vein_alchemist"],
    ["reverse_refinement", "逆炼归元", "恢复生命，并将一种玩家丹田灵气转为炉温。", "heal(10,self)+convert_opponent_qi_to_status(炉温)", "five_vein_alchemist"],
    ["furnace_explosion", "炸炉", "倒计时结束时造成极高全体伤害；Boss自身也损失生命。", "damage(危险,opponent)+damage(12,self)", "five_vein_alchemist"],
]

PROPOSED_ITEMS = [
    ["spring_powder", "回春散", "消耗品", "恢复10点生命。", "采药散修储物袋/普通掉落", "wandering_herbalist"],
    ["armor_break_talisman", "破甲符", "符箓", "造成8点金伤害，并削去6点护盾。", "百刃散人储物袋/商店", "hundred_blade_roamer"],
    ["mist_feather", "凝雾羽", "材料", "沾有水行雾气的长羽。", "雾苇鹤掉落", "mist_reed_crane"],
    ["stele_fragment", "负碑残片", "材料", "刻有残缺镇山纹的石片。", "负碑石猿掉落", "stoneback_ape"],
    ["golden_silk_gland", "金丝囊", "材料", "能够抽出细韧金丝的妖囊。", "金丝妖蛛掉落", "golden_thread_spider"],
    ["charcoal_core", "炭心妖丹", "材料", "内蕴躁动火气的妖丹。", "炭角火豕掉落", "char_horn_boar"],
    ["marsh_poison_sac", "泽毒囊", "材料", "水气与微毒纠缠的妖囊。", "泽母妖蟾掉落", "marsh_brood_mother"],
    ["five_vein_dregs", "五脉丹渣", "稀有材料", "五行药性互相冲突的丹炉残渣。", "五脉丹师掉落", "five_vein_alchemist"],
]


def load_data(filename: str, root_key: str):
    payload = json.loads((DATA_DIR / filename).read_text(encoding="utf-8"))
    return payload[root_key]


def zh(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float)):
        return value
    return ELEMENTS.get(
        str(value),
        RARITIES.get(str(value), UNIT_TYPES.get(str(value), TARGETS.get(str(value), str(value)))),
    )


def compact(value, level: int = 0) -> str:
    if value in (None, {}, []):
        return ""
    if isinstance(value, dict):
        return "；".join(f"{ELEMENTS.get(str(k), k)}={compact(v, level + 1)}" for k, v in value.items())
    if isinstance(value, list):
        sep = "\n" if any(isinstance(v, dict) for v in value) else "、"
        return sep.join(f"{i + 1}. {compact(v, level + 1)}" if isinstance(v, dict) else str(zh(v)) for i, v in enumerate(value))
    return str(zh(value))


def ids(value) -> str:
    return "、".join(map(str, value or []))


def effect_summary(effects) -> str:
    type_names = {
        "damage": "伤害", "shield": "护盾", "heal": "治疗", "buff": "强化",
        "add_status": "获得状态", "consume_status": "消耗状态", "add_sea_qi": "增加丹田灵气",
        "add_core_qi": "增加灵枢灵气", "consume_item": "消耗物品", "modifier": "规则修正",
        "trigger_effect": "触发效果", "learn_spell": "学习神通", "learn_method": "学习功法",
    }
    lines = []
    for i, effect in enumerate(effects or [], 1):
        if not isinstance(effect, dict):
            lines.append(str(effect))
            continue
        parts = [type_names.get(effect.get("type"), effect.get("type", "效果"))]
        if effect.get("attribute"):
            parts.append(ELEMENTS.get(effect["attribute"], effect["attribute"]))
        if effect.get("status"):
            parts.append(str(effect["status"]))
        if effect.get("value") is not None:
            parts.append(str(effect["value"]))
        if effect.get("target"):
            parts.append(f"→{TARGETS.get(effect['target'], effect['target'])}")
        extras = {k: v for k, v in effect.items() if k not in {"type", "attribute", "status", "value", "target"}}
        if extras:
            parts.append(f"[{compact(extras)}]")
        lines.append(f"{i}. " + " ".join(parts))
    return "\n".join(lines)


def action_summary(policy) -> str:
    if not policy:
        return ""
    lines = [f"策略：{policy.get('type', '')}"]
    for i, action in enumerate(policy.get("actions", []), 1):
        lines.append(f"{i}. {action.get('source_type', '')}:{action.get('source_id', '')}")
    return "\n".join(lines)


def create_sheet(wb: Workbook, name: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.append(headers)
    for row in rows:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="254E45")
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFF4D6")
    body_font = Font(name="Microsoft YaHei", size=10, color="2A2A2A")
    alt_fill = PatternFill("solid", fgColor="EFF6F1")
    line = Side(style="thin", color="D4C8A8")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="B78A3D"))
    ws.row_dimensions[1].height = 30

    for r in range(2, ws.max_row + 1):
        estimated_lines = 1
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=line)
            if r % 2 == 0:
                cell.fill = alt_fill
            if cell.value is not None:
                text = str(cell.value)
                col_width = widths[c - 1] if widths and c <= len(widths) else 16
                explicit = text.count("\n") + 1
                wrapped = math.ceil(max((len(part) for part in text.split("\n")), default=0) / max(col_width * 0.9, 1))
                estimated_lines = max(estimated_lines, explicit, wrapped)
        ws.row_dimensions[r].height = min(max(30, 18 * estimated_lines + 8), 110)

    if ws.max_row >= 2 and ws.max_column:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=f"T_{len(wb.sheetnames):02d}_{''.join(ch for ch in name if ch.isalnum())}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=False, showColumnStripes=False)
        ws.add_table(table)

    for idx in range(1, ws.max_column + 1):
        width = widths[idx - 1] if widths and idx <= len(widths) else 16
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 48)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    return ws


def main():
    enemies = load_data("enemies.json", "enemies")
    abilities = load_data("abilities.json", "abilities")
    spells = load_data("spells.json", "spells")
    items = load_data("items.json", "items")
    methods = load_data("cultivation_methods.json", "cultivation_methods")
    relics = load_data("relics.json", "relics")
    statuses = load_data("statuses.json", "statuses")
    battle_setups = load_data("battle_setups.json", "battle_setups")
    encounter_pools = load_data("encounter_pools.json", "encounter_pools")
    reward_tables = load_data("reward_tables.json", "reward_tables")
    node_contents = load_data("map_node_contents.json", "node_contents")
    node_offers = load_data("node_offer_tables.json", "node_offer_tables")
    qi_types = load_data("qi.json", "qi_types")
    spirit_cores = load_data("spirit_cores.json", "spirit_cores")
    player_templates = load_data("player_templates.json", "player_templates")
    recipes = load_data("processing_recipes.json", "processing_recipes")
    maps = load_data("map_nodes.json", "maps")
    map_generation = load_data("map_generation.json", "map_generation")
    breakthrough_rules = load_data("breakthrough_rules.json", "breakthrough_rules")
    attr = json.loads((DATA_DIR / "attribute_rules.json").read_text(encoding="utf-8"))

    wb = Workbook()
    wb.remove(wb.active)

    resource_counts = [
        ("敌人", len(enemies), "enemies.json"), ("行动能力", len(abilities), "abilities.json"),
        ("神通", len(spells), "spells.json"), ("功法", len(methods), "cultivation_methods.json"),
        ("法宝", len(relics), "relics.json"), ("物品", len(items), "items.json"),
        ("状态", len(statuses), "statuses.json"), ("战斗配置", len(battle_setups), "battle_setups.json"),
        ("遭遇池", len(encounter_pools), "encounter_pools.json"), ("奖励表", len(reward_tables), "reward_tables.json"),
        ("节点内容", len(node_contents), "map_node_contents.json"), ("节点候选", len(node_offers), "node_offer_tables.json"),
    ]
    overview = create_sheet(wb, "总览", ["资源类别", "当前数量", "来源文件", "策划用途"], [
        [name, count, src, {
            "敌人": "配置单位、行动来源与奖励引用", "行动能力": "敌人与通用行为效果", "神通": "玩家构筑卡池",
            "功法": "初始路线与规则修正", "法宝": "长期被动规则组件", "物品": "消耗品与材料",
            "状态": "战斗状态定义", "战斗配置": "具体战斗编队", "遭遇池": "随机遭遇来源",
            "奖励表": "战斗与节点奖励", "节点内容": "事件、修炼、休息等内容", "节点候选": "节点内随机候选池",
        }.get(name, "")]
        for name, count, src in resource_counts
    ], [16, 12, 28, 38])
    overview.insert_rows(1, 3)
    overview.merge_cells("A1:D1")
    overview["A1"] = "现有资源设计总表"
    overview["A1"].font = Font(name="Microsoft YaHei", size=18, bold=True, color="FFF4D6")
    overview["A1"].fill = PatternFill("solid", fgColor="183A34")
    overview["A1"].alignment = Alignment(horizontal="center", vertical="center")
    overview.row_dimensions[1].height = 38
    overview.merge_cells("A2:D2")
    overview["A2"] = "数据来自 data/*.json；中文列用于策划阅读，ID 与引用列用于回写数据时核对。"
    overview["A2"].alignment = Alignment(horizontal="left", vertical="center")
    overview["A2"].font = Font(name="Microsoft YaHei", italic=True, color="52645F")
    overview.freeze_panes = "A5"
    # The title rows sit above the table; rebuild the table reference after insertion.
    for table_name in list(overview.tables.keys()):
        del overview.tables[table_name]
    overview_table = Table(displayName="T_Overview", ref=f"A4:D{overview.max_row}")
    overview_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=False, showColumnStripes=False)
    overview.add_table(overview_table)

    create_sheet(wb, "敌人表", ["id", "名称", "单位类型", "最大生命", "属性倾向", "神通ID", "能力ID", "储物袋", "行动逻辑", "奖励表ID", "阶段预留", "姿态预留"], [
        [x.get("id"), x.get("name"), zh(x.get("unit_type")), x.get("max_hp"), compact(x.get("attributes", {}).get("tendencies")),
         ids(x.get("skills", {}).get("spell_ids")), ids(x.get("skills", {}).get("ability_ids")), compact(x.get("inventory")),
         action_summary(x.get("action_policy")), x.get("reward_table_id", ""), compact(x.get("phases")), compact(x.get("stances"))]
        for x in enemies
    ], [24, 16, 12, 10, 16, 24, 30, 18, 38, 26, 24, 24])

    create_sheet(wb, "技能表", ["id", "名称", "行动说明", "效果摘要", "原始效果字段"], [
        [x.get("id"), x.get("name"), x.get("description"), effect_summary(x.get("effects")), compact(x.get("effects"))] for x in abilities
    ], [28, 18, 42, 42, 48])

    create_sheet(wb, "神通表", ["id", "名称", "神识占用", "品质", "流派标签", "灵气消耗", "施展条件", "卡面描述", "效果摘要", "原始效果字段"], [
        [x.get("id"), x.get("name"), x.get("sense_cost", 0), zh(x.get("rarity")), ids([ELEMENTS.get(t, t) for t in x.get("tags", [])]),
         compact(x.get("cost")), compact(x.get("requirements")), x.get("description"), effect_summary(x.get("effects")), compact(x.get("effects"))]
        for x in spells
    ], [28, 18, 10, 10, 24, 18, 24, 42, 42, 48])

    create_sheet(wb, "功法表", ["id", "名称", "描述", "五行倾向", "初始神通ID", "替换神通ID", "效果摘要", "原始效果字段"], [
        [x.get("id"), x.get("name"), x.get("description"), compact(x.get("element_affinity")), x.get("initial_spell_id"), ids(x.get("replaces_spell_ids")), effect_summary(x.get("effects")), compact(x.get("effects"))]
        for x in methods
    ], [24, 18, 42, 18, 26, 24, 42, 48])

    create_sheet(wb, "法宝表", ["id", "名称", "图标ID", "标签", "描述", "触发与效果", "原始效果字段"], [
        [x.get("id"), x.get("name"), x.get("icon"), ids([ELEMENTS.get(t, t) for t in x.get("tags", [])]), x.get("description"), effect_summary(x.get("effects")), compact(x.get("effects"))]
        for x in relics
    ], [28, 18, 26, 24, 42, 44, 48])

    create_sheet(wb, "物品表", ["id", "名称", "类型", "标签", "描述", "战斗使用效果", "未来处理接口"], [
        [x.get("id"), x.get("name"), x.get("type"), ids([ELEMENTS.get(t, t) for t in x.get("tags", [])]), x.get("description"), effect_summary(x.get("combat_use", {}).get("effects")), compact(x.get("future_processing"))]
        for x in items
    ], [28, 18, 14, 22, 42, 42, 36])

    create_sheet(wb, "状态表", ["id", "名称", "描述", "叠加规则", "持续时间", "效果字段", "显示配置"], [
        [x.get("id"), x.get("name"), x.get("description"), compact(x.get("stacking")), x.get("duration", ""), compact(x.get("effects")), compact(x.get("display"))]
        for x in statuses
    ], [26, 18, 42, 20, 12, 44, 30])

    create_sheet(wb, "战斗配置", ["id", "名称", "遭遇类型", "敌人ID", "遭遇池ID", "奖励表ID", "场景/背景", "其他字段"], [
        [x.get("id"), x.get("name", ""), x.get("type", x.get("encounter_type", "")), ids(x.get("enemy_ids")), x.get("encounter_pool_id", ""), x.get("reward_table_id", ""), x.get("background", ""), compact({k:v for k,v in x.items() if k not in {"id","name","type","encounter_type","enemy_ids","encounter_pool_id","reward_table_id","background"}})]
        for x in battle_setups
    ], [28, 20, 14, 36, 28, 28, 24, 42])

    create_sheet(wb, "遭遇池", ["id", "类型", "敌人ID", "战斗配置ID", "权重/规则", "其他字段"], [
        [x.get("id"), x.get("type"), ids(x.get("enemy_ids")), ids(x.get("battle_setup_ids")), compact(x.get("weights", x.get("rules"))), compact({k:v for k,v in x.items() if k not in {"id","type","enemy_ids","battle_setup_ids","weights","rules"}})]
        for x in encounter_pools
    ], [30, 14, 42, 36, 30, 42])

    create_sheet(wb, "奖励表", ["id", "模式", "固定奖励", "神通候选", "灵气候选", "法宝候选", "其他奖励字段"], [
        [x.get("id"), x.get("mode", x.get("type", "")), compact(x.get("fixed_rewards", x.get("fixed"))), compact(x.get("spell_choices", x.get("spell_pool"))), compact(x.get("qi_choices", x.get("qi_pool"))), compact(x.get("relic_choices", x.get("relic_pool"))), compact({k:v for k,v in x.items() if k not in {"id","mode","type","fixed_rewards","fixed","spell_choices","spell_pool","qi_choices","qi_pool","relic_choices","relic_pool"}})]
        for x in reward_tables
    ], [30, 14, 46, 38, 34, 34, 48])

    create_sheet(wb, "节点内容", ["id", "名称", "节点类型", "描述", "选项/结果", "战斗/奖励引用", "其他字段"], [
        [x.get("id"), x.get("name"), x.get("type"), x.get("description", ""), compact(x.get("options", x.get("choices"))), compact({k:v for k,v in x.items() if "pool" in k or "battle" in k or "reward" in k}), compact({k:v for k,v in x.items() if k not in {"id","name","type","description","options","choices"} and "pool" not in k and "battle" not in k and "reward" not in k})]
        for x in node_contents
    ], [28, 20, 14, 42, 48, 36, 42])

    create_sheet(wb, "节点候选池", ["id", "节点类型", "候选内容", "权重", "条件", "其他字段"], [
        [x.get("id"), x.get("type", x.get("node_type", "")), compact(x.get("offers", x.get("entries", x.get("candidates")))), compact(x.get("weights")), compact(x.get("conditions")), compact({k:v for k,v in x.items() if k not in {"id","type","node_type","offers","entries","candidates","weights","conditions"}})]
        for x in node_offers
    ], [30, 14, 48, 24, 32, 42])

    create_sheet(wb, "扩展神通", ["id", "名称", "五行", "神识", "品质", "灵气消耗", "效果", "构筑作用", "建议来源", "状态"], [
        row + ["设计候选"] for row in PROPOSED_SPELLS
    ], [28, 18, 10, 10, 10, 16, 48, 24, 24, 12])

    create_sheet(wb, "五行法宝旧草案", ["id", "名称", "五行", "实际效果", "所需通用接口", "建议来源", "状态"], [
        row + ["暂缓：后续不再以五行绑定为主"] for row in PROPOSED_RELICS
    ], [28, 18, 10, 48, 34, 24, 12])

    create_sheet(wb, "扩展敌人", ["id", "名称", "级别", "单位类型", "属性倾向", "生命压力", "攻击压力", "配套能力", "行动循环", "战斗定位", "奖励表ID", "状态"], [
        row + ["设计候选"] for row in PROPOSED_ENEMIES
    ], [28, 18, 12, 12, 12, 12, 12, 38, 48, 42, 30, 12])

    create_sheet(wb, "扩展能力", ["id", "名称", "意图描述", "效果表达", "所属敌人ID", "状态"], [
        row + ["设计候选"] for row in PROPOSED_ABILITIES
    ], [32, 20, 44, 48, 30, 12])

    create_sheet(wb, "扩展物品", ["id", "名称", "类型", "效果/描述", "来源", "配套敌人ID", "状态"], [
        row + ["设计候选"] for row in PROPOSED_ITEMS
    ], [28, 18, 14, 44, 30, 30, 12])

    proposed_rewards = []
    material_by_enemy = {
        "mist_reed_crane": "mist_feather", "stoneback_ape": "stele_fragment",
        "golden_thread_spider": "golden_silk_gland", "char_horn_boar": "charcoal_core",
        "wandering_herbalist": "spring_powder", "grave_moss_wraith": "阴苔样本（待建物品）",
        "hundred_blade_roamer": "armor_break_talisman", "marsh_brood_mother": "marsh_poison_sac",
        "furnace_heart_golem": "炉心碎片（待建物品）", "five_vein_alchemist": "five_vein_dregs",
    }
    for enemy in PROPOSED_ENEMIES:
        enemy_id, enemy_name, tier = enemy[0], enemy[1], enemy[2]
        stones = "灵石×12~18" if tier == "普通" else ("灵石×35~50" if tier == "精英" else "灵石×90~120")
        special = "神通三选一 + 灵气三选一" if tier == "普通" else ("进阶神通三选一 + 法宝三选一" if tier == "精英" else "稀有神通三选一 + 稀有法宝三选一")
        proposed_rewards.append([
            f"{enemy_id}_rewards", enemy_name, tier, f"{stones}；{material_by_enemy[enemy_id]}×1", special,
            "候选按五行权重生成，但至少保留1个跨流派选项。", "设计候选",
        ])
    create_sheet(wb, "扩展奖励", ["奖励表ID", "对应敌人", "级别", "固定战利品", "选择奖励", "生成规则", "状态"],
                 proposed_rewards, [34, 20, 12, 42, 42, 46, 12])

    create_sheet(wb, "扩展接口词典", ["接口/状态", "类别", "用途", "首个使用内容", "建议实现边界", "状态"], [
        ["shield_break", "通用效果", "削减目标护盾，不溢出为生命伤害。", "破甲符", "EffectExecutor 通用效果", "待实现"],
        ["summon", "通用效果", "在空位生成指定单位并纳入统一 Unit 流程。", "泽母妖蟾", "BattleManager 通用召唤接口", "待实现"],
        ["cycle_attribute", "通用效果", "按数据数组轮转当前属性标签。", "五脉丹师", "状态字段驱动，不绑定敌人ID", "待实现"],
        ["convert_opponent_qi_to_status", "通用效果", "移除目标指定/随机丹田灵气并转为施法者状态。", "五脉丹师", "资源转化通用接口", "待实现"],
        ["错脉", "状态", "下一次凝聚时改变五行权重，随后移除。", "错脉散", "只影响丹田生成，不修改长期灵枢", "待实现"],
        ["湿毒", "状态", "回合结束损失少量生命；获得护盾时减少1层。", "湿毒吐息", "可净化、可衰减", "待实现"],
        ["炉温", "状态", "强化炉心石傀的部分能力，泄焰时清空。", "炉心石傀", "通用数值层数状态", "现有状态系统可表达"],
        ["焦躁", "状态", "强化下一次指定攻击并被其消耗。", "炭角火豕", "通用数值层数状态", "现有状态系统可表达"],
        ["藏锋", "状态", "记录下一次多段攻击次数。", "百刃散人", "通用数值层数状态", "现有状态系统可表达"],
        ["回合计数触发", "法宝触发器", "统计本回合某标签神通或灵气消耗次数。", "淬锋匣、三足火炉", "RelicManager 通用 trigger_counter", "待实现"],
    ], [32, 16, 46, 28, 46, 16])

    create_sheet(wb, "基础配置", ["配置类别", "id/键", "名称", "主要字段", "完整配置"],
        [["灵气类型", x.get("id"), x.get("name"), compact({k:v for k,v in x.items() if k not in {"id","name"}}), compact(x)] for x in qi_types]
        + [["灵枢模板", x.get("id"), x.get("name", ""), compact({k:v for k,v in x.items() if k not in {"id","name"}}), compact(x)] for x in spirit_cores]
        + [["玩家模板", x.get("id"), x.get("name", ""), compact({k:v for k,v in x.items() if k not in {"id","name"}}), compact(x)] for x in player_templates]
        + [["炼化配方", x.get("id"), x.get("name", ""), compact({k:v for k,v in x.items() if k not in {"id","name"}}), compact(x)] for x in recipes]
        + [["地图", x.get("id"), x.get("name", ""), compact({k:v for k,v in x.items() if k not in {"id","name"}}), compact(x)] for x in maps]
        + [["地图生成", str(k), "", compact(v), compact(v)] for k,v in (map_generation.items() if isinstance(map_generation,dict) else [(x.get("id",i),x) for i,x in enumerate(map_generation)])]
        + [["突破规则", x.get("id", i) if isinstance(x,dict) else i, x.get("name", "") if isinstance(x,dict) else "", compact(x), compact(x)] for i,x in enumerate(breakthrough_rules if isinstance(breakthrough_rules,list) else [breakthrough_rules])]
        + [["属性规则", "default_multiplier", "默认倍率", attr.get("default_multiplier"), compact(attr.get("default_multiplier"))], ["属性规则", "relationships", "五行关系", compact(attr.get("relationships")), compact(attr.get("relationships"))]],
        [16, 28, 20, 48, 48])

    # Apply consistent document properties and page setup.
    wb.properties.title = "现有资源设计表"
    wb.properties.subject = "灵枢修仙 Roguelike 当前数据资源汇总"
    wb.properties.creator = "Codex"
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_view.zoomScale = 85

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_BOOK)
    shutil.copy2(OUTPUT_BOOK, SOURCE_BOOK)
    print(OUTPUT_BOOK)
    print(SOURCE_BOOK)


if __name__ == "__main__":
    main()
