from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "设计表.xlsx"
OUT_DIR = ROOT / "设计表_CSV"
ZIP_BASE = ROOT / "设计表_CSV"


def safe_name(name: str) -> str:
    return "".join("_" if ch in '<>:"/\\|?*' else ch for ch in name).strip() or "Sheet"


def main() -> None:
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / ".gdignore").write_text("", encoding="utf-8")

    workbook = load_workbook(BOOK, data_only=False, read_only=True)
    manifest_rows = [["文件", "内容", "数据行数"]]

    for index, sheet in enumerate(workbook.worksheets, 1):
        filename = f"{index:02d}_{safe_name(sheet.title)}.csv"
        path = OUT_DIR / filename
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            for row in sheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])
        manifest_rows.append([filename, sheet.title, sheet.max_row])

    with (OUT_DIR / "00_目录.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows(manifest_rows)

    archive = shutil.make_archive(str(ZIP_BASE), "zip", root_dir=OUT_DIR)
    print(f"CSV_DIR={OUT_DIR}")
    print(f"ZIP={archive}")
    print(f"SHEETS={len(workbook.worksheets)}")


if __name__ == "__main__":
    main()
